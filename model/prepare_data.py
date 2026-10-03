from __future__ import annotations

import csv
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path
from statistics import median

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "source" / "applied_statistics_985_211_2026_source.xlsx"
OUT = ROOT / "data" / "processed"
AUDIT = ROOT / "data" / "audit"

NATIONAL_LINES = {
    2017: {"A": 335, "B": 325},
    2018: {"A": 330, "B": 320},
    2019: {"A": 345, "B": 335},
    2020: {"A": 343, "B": 333},
    2021: {"A": 348, "B": 338},
    2022: {"A": 360, "B": 350},
    2023: {"A": 346, "B": 336},
    2024: {"A": 338, "B": 328},
    2025: {"A": 323, "B": 313},
    2026: {"A": 324, "B": 314},
}

NATIONAL_LINE_SOURCES = {
    2022: "https://www.moe.gov.cn/jyb_xwfb/gzdt_gzdt/s5987/202203/W020220311666880715945.pdf",
    2023: "https://www.moe.gov.cn/jyb_xwfb/gzdt_gzdt/s5987/202303/W020230310707949584507.pdf",
    2024: "https://www.moe.gov.cn/jyb_xwfb/gzdt_gzdt/s5987/202403/W020240312598657688341.pdf",
    2025: "https://www.moe.gov.cn/jyb_xwfb/gzdt_gzdt/s5987/202502/W020250224570852404019.pdf",
    2026: "https://www.moe.gov.cn/jyb_xwfb/gzdt_gzdt/s5987/202602/W020260228373276320816.pdf",
}

B_ZONE_PATTERN = re.compile("内蒙古|广西|海南|贵州|云南|西藏|甘肃|青海|宁夏|新疆")
URL_PATTERN = re.compile(r"https?://[^\s；;]+")


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def as_float(value):
    value = clean(value)
    if value is None:
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def as_int(value):
    result = as_float(value)
    return None if result is None else int(round(result))


def yes_no(value):
    return 1 if str(clean(value) or "").lower() in {"是", "yes", "true", "1"} else 0


def evidence_class(text):
    text = str(clean(text) or "")
    if not text:
        return "missing", 0.0
    if "官方" in text and "机构" not in text and "转载" not in text:
        return "official", 1.0
    if "官方" in text:
        return "official_secondary", 0.85
    if "机构" in text:
        return "third_party", 0.62
    if "国家线" in text:
        return "national_floor_only", 0.35
    return "unclassified", 0.5


def lower_tail_q10(admitted_min, admitted_median, admitted_count):
    """Estimate sample Q10 from min/median/count when candidate-level scores are absent.

    The sample minimum is treated as approximately the 1/(n+1) order statistic.
    We linearly interpolate toward the median. This is deliberately transparent and
    receives lower label weight than an exact candidate-level quantile.
    """
    if admitted_min is None:
        return None, "missing", 0.0
    if admitted_median is None or admitted_count is None or admitted_count < 2:
        return admitted_min, "minimum_only", 0.35
    p_min = 1.0 / (admitted_count + 1.0)
    if p_min >= 0.10:
        return admitted_min, "small_n_minimum", 0.48
    fraction = (0.10 - p_min) / (0.50 - p_min)
    value = admitted_min + fraction * (admitted_median - admitted_min)
    return round(value, 2), "order_stat_interpolation", 0.68


def worksheet_records(ws, header_row):
    headers = [clean(cell.value) for cell in ws[header_row]]
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(clean(value) is None for value in row):
            continue
        record = {headers[i]: clean(value) for i, value in enumerate(row) if i < len(headers) and headers[i]}
        records.append(record)
    return records


def write_csv(path, rows, fieldnames=None):
    rows = list(rows)
    if fieldnames is None:
        fieldnames = []
        seen = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    fieldnames.append(key)
                    seen.add(key)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    AUDIT.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)

    school_source = worksheet_records(workbook["院校总表"], 2)
    year_source = worksheet_records(workbook["年度线与录取"], 4)
    source_registry = worksheet_records(workbook["来源台账"], 4)
    gap_source = worksheet_records(workbook["字段缺口审计"], 4)

    schools = []
    school_map = {}
    for row in school_source:
        school = clean(row.get("学校"))
        if not school:
            continue
        location = str(clean(row.get("学院校区地址（到区）")) or "")
        zone = "B" if B_ZONE_PATTERN.search(location) else "A"
        item = {
            "school": school,
            "is_985": yes_no(row.get("985")),
            "is_211": yes_no(row.get("211")),
            "unit_2026": clean(row.get("2026招生学院/单位")),
            "study_mode": clean(row.get("学习方式")),
            "politics_subject": clean(row.get("政治")),
            "english_subject": clean(row.get("英语")),
            "math_subject": clean(row.get("数学/经综")),
            "major_subject": clean(row.get("专业课")),
            "official_scope": clean(row.get("官方范围说明")),
            "official_books": clean(row.get("官方指定教材")),
            "official_book_status": clean(row.get("官方教材状态")),
            "reexam_subjects": clean(row.get("复试科目")),
            "reexam_freshness": clean(row.get("复试信息新鲜度")),
            "location": clean(row.get("学院校区地址（到区）")),
            "city_center": clean(row.get("市中心地标")),
            "duration": clean(row.get("学制")),
            "tuition": clean(row.get("学费")),
            "tuition_freshness": clean(row.get("学费新鲜度")),
            "data_completeness": clean(row.get("资料完整度")),
            "verified_at": clean(row.get("核验日期")),
            "national_zone": zone,
            "source_row": int(row.get("__row__", 0) or 0),
        }
        schools.append(item)
        school_map[school] = item

    gaps = {clean(row.get("学校")): row for row in gap_source if clean(row.get("学校"))}
    periods = []
    issues = []
    seen_keys = Counter()

    for source_row_index, row in enumerate(year_source, start=5):
        school = clean(row.get("学校"))
        year = as_int(row.get("年份"))
        if not school or year is None:
            continue
        seen_keys[(school, year)] += 1
        meta = school_map.get(school, {})
        zone = meta.get("national_zone", "A")
        cutoff = as_float(row.get("选择数值"))
        admitted_count = as_int(row.get("拟录取人数"))
        admitted_min = as_float(row.get("拟录取最低"))
        admitted_median = as_float(row.get("拟录取中位"))
        admitted_max = as_float(row.get("拟录取最高"))
        retest_count = as_int(row.get("复试人数"))
        ratio_reported = as_float(row.get("复录比"))
        ratio_calculated = None
        if admitted_count and retest_count is not None:
            ratio_calculated = retest_count / admitted_count

        evidence, cutoff_weight = evidence_class(row.get("线证据等级"))
        q10, q10_method, q10_method_weight = lower_tail_q10(
            admitted_min, admitted_median, admitted_count
        )
        admit_note = str(clean(row.get("录取来源/备注")) or "")
        admission_official = bool(re.search(r"(?:\.edu\.cn|\.ac\.cn|yz\.chsi\.com\.cn)", admit_note))
        q10_weight = q10_method_weight * (1.0 if admission_official else 0.72)
        national_line = NATIONAL_LINES.get(year, {}).get(zone)
        margin_cutoff = None if cutoff is None or national_line is None else cutoff - national_line
        margin_q10 = None if q10 is None or national_line is None else q10 - national_line
        target = q10 if q10 is not None else cutoff
        target_kind = "q10_proxy" if q10 is not None else "cutoff_fallback"
        target_weight = q10_weight if q10 is not None else cutoff_weight * 0.55

        row_issues = []
        if school not in school_map:
            row_issues.append("school_missing_from_master")
        if cutoff is not None and not (250 <= cutoff <= 500):
            row_issues.append("cutoff_out_of_range")
        ordered = [v for v in (admitted_min, admitted_median, admitted_max) if v is not None]
        if len(ordered) >= 2 and ordered != sorted(ordered):
            row_issues.append("admission_summary_order_violation")
        if cutoff is not None and admitted_min is not None and admitted_min + 1e-9 < cutoff:
            row_issues.append("admitted_min_below_selected_cutoff")
        if ratio_reported is not None and ratio_calculated is not None and abs(ratio_reported - ratio_calculated) > 0.051:
            row_issues.append("retest_ratio_mismatch")
        if cutoff is not None and not clean(row.get("分数线来源")):
            row_issues.append("cutoff_missing_source")
        if any(v is not None for v in (admitted_min, admitted_median, admitted_max)) and not admit_note:
            row_issues.append("admission_summary_missing_source")
        if "专项" in admit_note and "排除" not in admit_note:
            row_issues.append("special_plan_exclusion_unclear")

        for code in row_issues:
            issues.append({
                "school": school,
                "year": year,
                "source_sheet": "年度线与录取",
                "source_row": source_row_index,
                "issue_code": code,
                "severity": "high" if code in {"admission_summary_order_violation", "retest_ratio_mismatch"} else "medium",
                "review_status": "待人工复核",
            })

        periods.append({
            "school": school,
            "year": year,
            "national_zone": zone,
            "national_line": national_line,
            "cutoff": cutoff,
            "cutoff_text": clean(row.get("复试线文本")),
            "cutoff_evidence": clean(row.get("线证据等级")),
            "cutoff_evidence_class": evidence,
            "cutoff_weight": round(cutoff_weight, 3),
            "cutoff_detail": clean(row.get("多学院/方向明细")),
            "cutoff_source": clean(row.get("分数线来源")),
            "admitted_count": admitted_count,
            "admitted_min": admitted_min,
            "admitted_median": admitted_median,
            "admitted_max": admitted_max,
            "retest_count": retest_count,
            "retest_ratio_reported": ratio_reported,
            "retest_ratio_calculated": None if ratio_calculated is None else round(ratio_calculated, 4),
            "q10_value": q10,
            "q10_method": q10_method,
            "q10_weight": round(q10_weight, 3),
            "target_value": target,
            "target_kind": target_kind,
            "target_weight": round(target_weight, 3),
            "margin_cutoff": margin_cutoff,
            "margin_q10": margin_q10,
            "admission_official_url_present": int(admission_official),
            "admission_source_note": clean(row.get("录取来源/备注")),
            "issue_count": len(row_issues),
            "issue_codes": ";".join(row_issues) if row_issues else None,
            "source_sheet": "年度线与录取",
            "source_row": source_row_index,
        })

    for (school, year), count in seen_keys.items():
        if count > 1:
            issues.append({
                "school": school,
                "year": year,
                "source_sheet": "年度线与录取",
                "source_row": None,
                "issue_code": "duplicate_school_year",
                "severity": "high",
                "review_status": "待人工复核",
            })

    by_school = defaultdict(list)
    for row in periods:
        by_school[row["school"]].append(row)
    for school_rows in by_school.values():
        school_rows.sort(key=lambda row: row["year"])
        prior_margins = []
        previous = None
        previous_previous = None
        for row in school_rows:
            row["lag1_cutoff"] = previous.get("cutoff") if previous else None
            row["lag1_q10"] = previous.get("q10_value") if previous else None
            row["lag1_margin_q10"] = previous.get("margin_q10") if previous else None
            row["lag1_admitted_count"] = previous.get("admitted_count") if previous else None
            row["lag1_cutoff_change"] = (
                None
                if previous is None
                or previous_previous is None
                or previous.get("cutoff") is None
                or previous_previous.get("cutoff") is None
                else previous["cutoff"] - previous_previous["cutoff"]
            )
            row["lag1_margin_change"] = (
                None
                if previous is None
                or previous_previous is None
                or previous.get("margin_q10") is None
                or previous_previous.get("margin_q10") is None
                else previous["margin_q10"] - previous_previous["margin_q10"]
            )
            row["trailing_margin_median"] = median(prior_margins) if prior_margins else None
            history_before_previous = prior_margins[:-1]
            if previous and previous.get("margin_q10") is not None and history_before_previous:
                center = median(history_before_previous)
                absolute = [abs(v - center) for v in history_before_previous]
                scale = median(absolute) * 1.4826 if len(absolute) >= 2 else 12.0
                if scale < 4:
                    scale = 4.0
                row["lag1_surprise_z"] = (previous["margin_q10"] - center) / scale
            else:
                row["lag1_surprise_z"] = None
            if row.get("margin_q10") is not None:
                prior_margins.append(row["margin_q10"])
            previous_previous = previous
            previous = row

    peer_groups = defaultdict(list)
    broad_peer_groups = defaultdict(list)
    for row in periods:
        meta = school_map.get(row["school"], {})
        trailing = row.get("trailing_margin_median")
        band = None if trailing is None else int(round(trailing / 20.0))
        peer_key = (row["year"], meta.get("is_985", 0), row["national_zone"], band)
        broad_key = (row["year"], meta.get("is_985", 0))
        if row.get("lag1_surprise_z") is not None:
            peer_groups[peer_key].append((row["school"], row["lag1_surprise_z"]))
            broad_peer_groups[broad_key].append((row["school"], row["lag1_surprise_z"]))

    for row in periods:
        meta = school_map.get(row["school"], {})
        trailing = row.get("trailing_margin_median")
        band = None if trailing is None else int(round(trailing / 20.0))
        peer_key = (row["year"], meta.get("is_985", 0), row["national_zone"], band)
        broad_key = (row["year"], meta.get("is_985", 0))
        candidates = [value for school, value in peer_groups.get(peer_key, []) if school != row["school"]]
        if len(candidates) < 2:
            candidates = [value for school, value in broad_peer_groups.get(broad_key, []) if school != row["school"]]
        row["peer_lag1_surprise_mean"] = (
            round(sum(candidates) / len(candidates), 4) if candidates else None
        )

    source_rows = []
    for index, row in enumerate(source_registry, start=5):
        url_text = str(clean(row.get("URL")) or "")
        urls = URL_PATTERN.findall(url_text)
        source_rows.append({
            "source_name": clean(row.get("来源名称")),
            "evidence_level": clean(row.get("等级")),
            "covered_fields": clean(row.get("覆盖字段")),
            "url": url_text or None,
            "url_count": len(urls),
            "notes": clean(row.get("备注")),
            "source_sheet": "来源台账",
            "source_row": index,
        })

    gap_rows = []
    for school, row in gaps.items():
        gap_rows.append({
            "school": school,
            "missing_cutoff_years": clean(row.get("复试线数值缺失年份")),
            "weak_cutoff_years": clean(row.get("复试线弱证据年份")),
            "missing_admission_years": clean(row.get("拟录取统计缺失年份")),
            "latest_retest_ratio": clean(row.get("最新复录比")),
            "official_books_status": clean(row.get("官方教材")),
            "tuition_status": clean(row.get("学费")),
            "exam_evidence_status": clean(row.get("真题证据")),
            "priority": clean(row.get("优先级")),
            "next_official_channel": clean(row.get("下一官方检索渠道")),
            "audit_note": clean(row.get("审计备注")),
        })

    write_csv(OUT / "schools.csv", schools)
    write_csv(OUT / "program_year.csv", periods)
    write_csv(OUT / "source_registry.csv", source_rows)
    write_csv(OUT / "review_queue.csv", issues)
    write_csv(OUT / "field_gaps.csv", gap_rows)

    exact_q10_count = sum(1 for row in periods if row["q10_method"] == "exact_candidate_scores")
    report = {
        "source_file": SOURCE.name,
        "school_count": len(schools),
        "program_year_rows": len(periods),
        "year_range": [min(row["year"] for row in periods), max(row["year"] for row in periods)],
        "unique_school_years": len(seen_keys),
        "duplicate_school_years": sum(1 for count in seen_keys.values() if count > 1),
        "cutoff_available": sum(row["cutoff"] is not None for row in periods),
        "cutoff_official_or_official_secondary": sum(row["cutoff_evidence_class"].startswith("official") for row in periods),
        "admission_min_available": sum(row["admitted_min"] is not None for row in periods),
        "admission_median_available": sum(row["admitted_median"] is not None for row in periods),
        "q10_proxy_available": sum(row["q10_value"] is not None for row in periods),
        "q10_exact_available": exact_q10_count,
        "review_issue_count": len(issues),
        "review_issue_types": dict(Counter(row["issue_code"] for row in issues)),
        "important_limitations": [
            "The source workbook does not contain candidate-level Q10 values. Q10 is transparently estimated from minimum, median, and admitted count where available.",
            "Historical realized admitted count is not identical to the plan known before registration and must not be used as a leakage-free current-year seat feature.",
            "Rows marked third-party remain lower-weight observations pending official verification.",
            "Multiple colleges or directions are aggregated conservatively at school level when the source workbook uses the highest ordinary full-time cutoff.",
        ],
    }
    (AUDIT / "source_audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUT / "national_lines.json").write_text(
        json.dumps({"values": NATIONAL_LINES, "sources": NATIONAL_LINE_SOURCES}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
