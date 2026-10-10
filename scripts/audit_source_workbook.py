from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl


URL_PATTERN = re.compile(r"https?://[^\s，。；、)）\]\}]+", re.IGNORECASE)
SOURCE_TERMS = ("来源", "链接", "网址", "url", "官网", "证据", "出处", "备注")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audit provenance fields in the source workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def cell_record(cell: openpyxl.cell.cell.Cell) -> dict | None:
    value = cell.value
    hyperlink = cell.hyperlink.target if cell.hyperlink else None
    comment = None
    if cell.comment:
        comment = {"author": cell.comment.author, "text": cell.comment.text}
    urls = URL_PATTERN.findall(str(value)) if value is not None else []
    if value is None and hyperlink is None and comment is None:
        return None
    return {
        "coordinate": cell.coordinate,
        "value": value,
        "data_type": cell.data_type,
        "hyperlink": hyperlink,
        "comment": comment,
        "urls_in_value": urls,
    }


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_dir = args.output.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    workbook = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    audit = {
        "workbook": str(workbook_path),
        "sheets": [],
        "totals": {},
    }
    all_hyperlinks = []
    all_urls = []
    comment_count = 0
    formula_count = 0

    for sheet in workbook.worksheets:
        rows = []
        source_header_cells = []
        hyperlink_count = 0
        sheet_comment_count = 0
        sheet_formula_count = 0
        for row in sheet.iter_rows():
            records = []
            for cell in row:
                record = cell_record(cell)
                if record is None:
                    continue
                records.append(record)
                text_value = str(record["value"] or "")
                if any(term in text_value.lower() for term in SOURCE_TERMS):
                    source_header_cells.append({"coordinate": cell.coordinate, "value": cell.value})
                if record["hyperlink"]:
                    hyperlink_count += 1
                    all_hyperlinks.append(
                        {"sheet": sheet.title, "cell": cell.coordinate, "target": record["hyperlink"]}
                    )
                for url in record["urls_in_value"]:
                    all_urls.append({"sheet": sheet.title, "cell": cell.coordinate, "target": url})
                if record["comment"]:
                    sheet_comment_count += 1
                if cell.data_type == "f":
                    sheet_formula_count += 1
            if records:
                rows.append({"row": row[0].row, "cells": records})

        comment_count += sheet_comment_count
        formula_count += sheet_formula_count
        audit["sheets"].append(
            {
                "title": sheet.title,
                "state": sheet.sheet_state,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
                "nonempty_rows": len(rows),
                "hyperlink_count": hyperlink_count,
                "comment_count": sheet_comment_count,
                "formula_count": sheet_formula_count,
                "source_header_cells": source_header_cells,
                "rows": rows,
            }
        )

    unique_targets = Counter(item["target"] for item in all_hyperlinks + all_urls)
    audit["totals"] = {
        "sheet_count": len(workbook.worksheets),
        "hyperlink_cells": len(all_hyperlinks),
        "urls_in_values": len(all_urls),
        "unique_url_targets": len(unique_targets),
        "comment_cells": comment_count,
        "formula_cells": formula_count,
    }
    audit["url_inventory"] = [
        {"target": target, "occurrences": count}
        for target, count in unique_targets.most_common()
    ]
    audit["hyperlink_cells"] = all_hyperlinks
    audit["url_cells"] = all_urls

    (output_dir / "workbook_source_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2, default=str),
        encoding="utf-8",
    )
    print(json.dumps(audit["totals"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
