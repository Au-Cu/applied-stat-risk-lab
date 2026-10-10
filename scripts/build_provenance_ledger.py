"""Build field-level provenance ledgers from the research source workbook.

The workbook is a researcher-compiled index, not itself the public evidence.
This script preserves the complete chain from a processed value back to the
workbook cell and then to every public URL recorded in that cell.  When a cell
contains several URLs the linkage is deliberately labelled ``cell_scope``;
the script does not pretend that the workbook identifies which URL supports
which individual value.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "source" / "applied_statistics_985_211_2026_source.xlsx"
DEFAULT_OUTPUT = ROOT / "data" / "processed"
DEFAULT_AUDIT = ROOT / "data" / "audit" / "provenance_coverage.json"

URL_PATTERN = re.compile(r"https?://[^\s<>'\"，。；、]+", re.IGNORECASE)
URL_TRAILING = ",;:!?)]}，。；、！？）】》"

YEAR_FIELD_GROUPS = {
    "cutoff": {
        "source_column": "分数线来源",
        "fields": ("选择数值", "复试线文本", "线证据等级", "多学院/方向明细"),
    },
    "admission": {
        "source_column": "录取来源/备注",
        "fields": (
            "拟录取人数",
            "拟录取最低",
            "拟录取中位",
            "拟录取最高",
            "复试人数",
            "复录比",
            "双非占比",
            "双非最低",
            "双非中位",
            "溢出率",
            "本科背景状态",
        ),
    },
}

MASTER_FIELD_GROUPS = {
    "exam_syllabus": {
        "source_column": "真题/大纲来源",
        "fields": (
            "政治",
            "英语",
            "数学/经综",
            "专业课",
            "官方范围说明",
            "官方指定教材",
            "官方教材状态",
            "真题证据等级",
            "最新官方大纲题型",
        ),
    }
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build normalized public-source provenance ledgers.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--audit", type=Path, default=DEFAULT_AUDIT)
    return parser.parse_args()


def sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def extract_urls(value) -> list[str]:
    urls = []
    for match in URL_PATTERN.findall(text(value)):
        url = match.rstrip(URL_TRAILING)
        if url and url not in urls:
            urls.append(url)
    return urls


def domain(url: str) -> str:
    try:
        return (urlsplit(url).hostname or "").lower()
    except ValueError:
        return ""


def classify_source(url: str) -> tuple[str, str, int]:
    """Return (source_type, authority_grade, authority_rank)."""

    host = domain(url)
    lower = url.lower()
    if not url:
        return "missing_public_url", "unresolved", 9
    if host.endswith(".edu.cn") or host.endswith(".ac.cn"):
        return "university_official", "primary_official", 1
    if host.endswith("moe.gov.cn") or host.endswith("chsi.com.cn"):
        return "government_official", "primary_official", 1
    if host == "admin.jiuxianzai.cn":
        return "mini_program_api", "secondary_structured", 2
    if host == "jiuxianzai.oss-cn-beijing.aliyuncs.com":
        return "mini_program_asset", "secondary_snapshot", 2
    if host == "mp.weixin.qq.com":
        return "wechat_article", "secondary_article", 3
    if lower.endswith((".png", ".jpg", ".jpeg", ".webp")):
        return "image_asset", "unclassified_snapshot", 4
    if lower.endswith(".pdf"):
        return "document_pdf", "unclassified_document", 4
    return "web_page", "unclassified_web", 4


def infer_artifact_role(url: str, group: str) -> str:
    host = domain(url)
    path = urlsplit(url).path.lower() if url else ""
    if host == "admin.jiuxianzai.cn" and "schooldetail" in path:
        return "mini_program_school_detail"
    if host == "jiuxianzai.oss-cn-beijing.aliyuncs.com" or path.endswith(
        (".png", ".jpg", ".jpeg", ".webp")
    ):
        return "roster_or_statistics_image" if group == "admission" else "cutoff_image"
    if path.endswith(".pdf"):
        return f"{group}_pdf"
    if host == "mp.weixin.qq.com":
        return f"{group}_wechat_article"
    return f"{group}_web_page"


def worksheet_headers(ws, header_row: int) -> tuple[list[str], dict[str, int]]:
    headers = [text(cell.value) for cell in ws[header_row]]
    return headers, {value: index + 1 for index, value in enumerate(headers) if value}


def write_csv(path: Path, rows: list[dict], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def build_field_rows(
    ws,
    *,
    header_row: int,
    groups: dict[str, dict],
    key_fields: tuple[str, ...],
    workbook_sha256: str,
) -> list[dict]:
    _, columns = worksheet_headers(ws, header_row)
    rows: list[dict] = []
    for row_index in range(header_row + 1, ws.max_row + 1):
        keys = {key: text(ws.cell(row_index, columns[key]).value) for key in key_fields}
        if not all(keys.values()):
            continue
        record_id = "|".join(keys[key] for key in key_fields)
        for group_name, group in groups.items():
            source_column = group["source_column"]
            source_cell = ws.cell(row_index, columns[source_column])
            source_text = text(source_cell.value)
            urls = extract_urls(source_text)
            sources = urls or [""]
            for field_name in group["fields"]:
                if field_name not in columns:
                    continue
                value_cell = ws.cell(row_index, columns[field_name])
                value = value_cell.value
                if value is None or text(value) == "":
                    continue
                for source_index, url in enumerate(sources, start=1):
                    source_type, authority_grade, authority_rank = classify_source(url)
                    source_id = hashlib.sha256(url.encode("utf-8")).hexdigest()[:16] if url else ""
                    rows.append(
                        {
                            **keys,
                            "record_id": record_id,
                            "field_group": group_name,
                            "field_name": field_name,
                            "field_value": value,
                            "source_url": url,
                            "source_id": source_id,
                            "source_index_in_cell": source_index if url else "",
                            "source_count_in_cell": len(urls),
                            "source_domain": domain(url),
                            "source_type": source_type,
                            "artifact_role": infer_artifact_role(url, group_name),
                            "authority_grade": authority_grade,
                            "authority_rank": authority_rank,
                            "linkage_precision": "cell_scope" if len(urls) != 1 else "direct_cell",
                            "public_url_status": "recorded" if url else "missing",
                            "source_text": source_text,
                            "workbook_sha256": workbook_sha256,
                            "source_sheet": ws.title,
                            "value_cell": value_cell.coordinate,
                            "source_cell": source_cell.coordinate,
                            "source_row": row_index,
                        }
                    )
    return rows


def build_registry_rows(ws, workbook_sha256: str) -> list[dict]:
    _, columns = worksheet_headers(ws, 4)
    rows = []
    for row_index in range(5, ws.max_row + 1):
        source_name = text(ws.cell(row_index, columns["来源名称"]).value)
        if not source_name:
            continue
        raw_url = text(ws.cell(row_index, columns["URL"]).value)
        urls = extract_urls(raw_url) or [""]
        for index, url in enumerate(urls, start=1):
            source_type, authority_grade, authority_rank = classify_source(url)
            rows.append(
                {
                    "source_name": source_name,
                    "declared_level": text(ws.cell(row_index, columns["等级"]).value),
                    "covered_fields": text(ws.cell(row_index, columns["覆盖字段"]).value),
                    "source_url": url,
                    "source_index_in_cell": index if url else "",
                    "source_count_in_cell": len([item for item in urls if item]),
                    "source_domain": domain(url),
                    "source_type": source_type,
                    "authority_grade": authority_grade,
                    "authority_rank": authority_rank,
                    "notes": text(ws.cell(row_index, columns["备注"]).value),
                    "workbook_sha256": workbook_sha256,
                    "source_sheet": ws.title,
                    "source_cell": ws.cell(row_index, columns["URL"]).coordinate,
                    "source_row": row_index,
                }
            )
    return rows


def coverage_summary(rows: list[dict], keys: tuple[str, ...]) -> dict:
    fields = defaultdict(lambda: {"urls": set(), "has_url": False, "authority": Counter()})
    records = defaultdict(lambda: {"fields": set(), "fields_with_url": set(), "urls": set()})
    for row in rows:
        field_key = tuple(row[key] for key in keys) + (row["field_name"], row["value_cell"])
        record_key = tuple(row[key] for key in keys)
        fields[field_key]["has_url"] |= bool(row["source_url"])
        if row["source_url"]:
            fields[field_key]["urls"].add(row["source_url"])
            fields[field_key]["authority"][row["authority_grade"]] += 1
            records[record_key]["fields_with_url"].add(row["field_name"])
            records[record_key]["urls"].add(row["source_url"])
        records[record_key]["fields"].add(row["field_name"])
    field_total = len(fields)
    field_with_url = sum(item["has_url"] for item in fields.values())
    return {
        "record_count": len(records),
        "field_value_count": field_total,
        "field_values_with_public_url": field_with_url,
        "field_url_coverage": round(field_with_url / field_total, 4) if field_total else None,
        "records_with_any_public_url": sum(bool(item["urls"]) for item in records.values()),
        "unique_public_urls": len({url for item in records.values() for url in item["urls"]}),
        "source_types": dict(Counter(row["source_type"] for row in rows if row["source_url"])),
        "authority_grades": dict(Counter(row["authority_grade"] for row in rows if row["source_url"])),
    }


def main() -> None:
    args = parse_args()
    workbook_path = args.workbook.resolve()
    output_dir = args.output_dir.resolve()
    audit_path = args.audit.resolve()
    workbook_sha256 = sha256_path(workbook_path)
    workbook = openpyxl.load_workbook(workbook_path, data_only=True, read_only=False)

    year_rows = build_field_rows(
        workbook["年度线与录取"],
        header_row=4,
        groups=YEAR_FIELD_GROUPS,
        key_fields=("学校", "年份"),
        workbook_sha256=workbook_sha256,
    )
    master_rows = build_field_rows(
        workbook["院校总表"],
        header_row=2,
        groups=MASTER_FIELD_GROUPS,
        key_fields=("学校",),
        workbook_sha256=workbook_sha256,
    )
    registry_rows = build_registry_rows(workbook["来源台账"], workbook_sha256)

    common_fields = [
        "record_id",
        "field_group",
        "field_name",
        "field_value",
        "source_url",
        "source_id",
        "source_index_in_cell",
        "source_count_in_cell",
        "source_domain",
        "source_type",
        "artifact_role",
        "authority_grade",
        "authority_rank",
        "linkage_precision",
        "public_url_status",
        "source_text",
        "workbook_sha256",
        "source_sheet",
        "value_cell",
        "source_cell",
        "source_row",
    ]
    write_csv(output_dir / "program_year_provenance.csv", year_rows, ["学校", "年份", *common_fields])
    write_csv(output_dir / "school_master_provenance.csv", master_rows, ["学校", *common_fields])
    write_csv(
        output_dir / "source_registry_urls.csv",
        registry_rows,
        [
            "source_name",
            "declared_level",
            "covered_fields",
            "source_url",
            "source_index_in_cell",
            "source_count_in_cell",
            "source_domain",
            "source_type",
            "authority_grade",
            "authority_rank",
            "notes",
            "workbook_sha256",
            "source_sheet",
            "source_cell",
            "source_row",
        ],
    )

    audit = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "workbook": str(workbook_path.relative_to(ROOT)).replace("\\", "/"),
        "workbook_sha256": workbook_sha256,
        "interpretation": (
            "The workbook is a compiled research index. Public URLs, not the workbook alone, "
            "are the underlying evidence. Multi-URL source cells are linked at cell scope."
        ),
        "program_year": coverage_summary(year_rows, ("学校", "年份")),
        "school_master_exam_syllabus": coverage_summary(master_rows, ("学校",)),
        "registry": {
            "rows": len(registry_rows),
            "entries_with_public_url": sum(bool(row["source_url"]) for row in registry_rows),
            "unique_public_urls": len({row["source_url"] for row in registry_rows if row["source_url"]}),
        },
    }
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(audit, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
